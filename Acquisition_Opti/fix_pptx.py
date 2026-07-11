"""
Correction ciblée : retire les images T2 mal placées sur slide 2,
les ajoute correctement sur slide 15.
"""
import os
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE_TYPE

THIS_DIR = os.path.dirname(os.path.abspath(__file__))
PPTX_PATH = os.path.join(THIS_DIR, "Figures", "DWI acquisition optimisation.pptx")

FIGS = {
    "fig_t2_optimum":  os.path.join(THIS_DIR, "Figures", "t2_sensitivity", "fig_t2_optimum.png"),
    "fig_t2_heatmaps": os.path.join(THIS_DIR, "Figures", "t2_sensitivity", "fig_t2_heatmaps.png"),
}

prs = Presentation(PPTX_PATH)
W = prs.slide_width.inches
H = prs.slide_height.inches

print(f"Total slides: {len(prs.slides)}")

# ── 1. Retirer les images de slide 2 (index 1) ────────────────────────────────
slide2 = prs.slides[1]
spTree = slide2.shapes._spTree
pics_to_remove = []
for shape in slide2.shapes:
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        pics_to_remove.append(shape._element)
        print(f"  Removing image from slide 2: {shape.name}")

for el in pics_to_remove:
    spTree.remove(el)

print(f"Removed {len(pics_to_remove)} images from slide 2")

# ── 2. Ajouter les images T2 sur slide 15 (index 14) ─────────────────────────
slide_t2 = prs.slides[14]
print(f"Slide 15 title: {slide_t2.shapes.title.text[:50]!r}")

# Vérifier qu'il n'y a pas déjà des images
existing_pics = [s for s in slide_t2.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
if existing_pics:
    print(f"  Slide 15 already has {len(existing_pics)} images — skipping")
else:
    H_img = 5.50
    w_opt  = H_img * 1715 / 3302   # ≈ 2.86"
    w_heat = H_img * 2475 / 3601   # ≈ 3.78"
    gap    = 0.40
    total_w = w_opt + gap + w_heat
    left_opt  = (W - total_w) / 2
    left_heat = left_opt + w_opt + gap
    top_img   = 1.65

    slide_t2.shapes.add_picture(FIGS["fig_t2_optimum"],
                                 Inches(left_opt), Inches(top_img),
                                 Inches(w_opt), Inches(H_img))
    slide_t2.shapes.add_picture(FIGS["fig_t2_heatmaps"],
                                 Inches(left_heat), Inches(top_img),
                                 Inches(w_heat), Inches(H_img))
    print("  T2 images added to slide 15 ✓")

# ── 3. Vérification finale ────────────────────────────────────────────────────
print("\nFinal slide structure:")
for i, slide in enumerate(prs.slides):
    n_pics = sum(1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    try:
        title = slide.shapes.title.text[:50].replace('\n',' ').replace('\x0b',' ')
    except:
        title = "(no title)"
    print(f"  Slide {i+1:2d}: {title!r}  [{n_pics} img]")

prs.save(PPTX_PATH)
print(f"\nSaved -> {PPTX_PATH}")
